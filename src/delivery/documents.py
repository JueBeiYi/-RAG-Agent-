from __future__ import annotations

import ast
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .inputs import COURSE_SECTIONS, DeliveryInputs, digest, is_demo, metric_rows, module_rows

BLACK = RGBColor(0, 0, 0)


def _font(run, size=11.5, *, bold=False, family="宋体"):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), family)


class Writer:
    def __init__(self, title: str, template: Path | None = None):
        self.doc = Document(template) if template else Document()
        if template:
            for element in list(self.doc._element.body):
                if element.tag != qn("w:sectPr"):
                    self.doc._element.body.remove(element)
        else:
            section = self.doc.sections[0]
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.3)
            section.left_margin = section.right_margin = Cm(2.5)
        self.template = template
        self.body_size = 14 if template else 11.5
        for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
            try:
                style = self.doc.styles[name]
            except KeyError:
                style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                style.base_style = self.doc.styles["Heading 2" if name == "Heading 3" else "Normal"]
            style.font.color.rgb = BLACK
            for element in list(style.element):
                if element.tag == qn("w:pPr"):
                    for border in list(element):
                        if border.tag == qn("w:pBdr"):
                            element.remove(border)
        self.title(title)

    def title(self, text):
        paragraph = self.doc.add_paragraph(style="Title")
        paragraph.paragraph_format.space_after = Pt(16)
        _font(paragraph.add_run(text), 22 if self.template else 24, bold=True, family="黑体")
        return paragraph

    def heading(self, text, level=1):
        paragraph = self.doc.add_paragraph(style=f"Heading {min(level, 3)}")
        _font(paragraph.add_run(text), 16 if level == 1 else 13, bold=True, family="黑体")
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def text(self, text, *, bold=False):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5 if self.template else 1.35
        paragraph.paragraph_format.space_after = Pt(7)
        _font(paragraph.add_run(str(text)), self.body_size, bold=bold, family="仿宋" if self.template else "宋体")
        return paragraph

    def table(self, headers, rows, widths=None):
        section = self.doc.sections[-1]
        available = (section.page_width - section.left_margin - section.right_margin) / 360000
        if widths:
            widths = [width * min(1, available / sum(widths)) for width in widths]
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        if widths:
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for cell, width in zip(table.rows[0].cells, widths):
                cell.width = Cm(width)
        for column, header in enumerate(headers):
            table.rows[0].cells[column].text = str(header)
        header_prop = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        header_prop.append(repeat)
        for values in rows:
            cells = table.add_row().cells
            for column, value in enumerate(values):
                cells[column].text = str(value)
                if widths:
                    cells[column].width = Cm(widths[column])
        for i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for j, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "203C50" if i == 0 else "EEF3F6" if i % 2 == 0 else "FFFFFF")
                properties.append(shading)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    border.set(qn("w:val"), "single"); border.set(qn("w:sz"), "4"); border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                properties.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), "90"); margin.set(qn("w:type"), "dxa"); margins.append(margin)
                properties.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.15
                    if j > 0 and len(cell.text) < 25:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        _font(run, 10.5, bold=i == 0)
                        if i == 0:
                            run.font.color.rgb = RGBColor(255, 255, 255)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return table

    def image(self, path: Path, caption: str):
        if not path.is_file():
            self.text("图示尚未提供：" + caption)
            return
        section = self.doc.sections[-1]
        width = section.page_width - section.left_margin - section.right_margin
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=width)
        caption_paragraph = self.doc.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(caption_paragraph.add_run(caption), 10)

    def save(self, path: Path):
        self.doc.core_properties.author = ""
        self.doc.core_properties.last_modified_by = ""
        self.doc.core_properties.created = datetime.now(timezone.utc)
        self.doc.core_properties.modified = datetime.now(timezone.utc)
        settings = self.doc.settings.element
        update = settings.find(qn("w:updateFields"))
        if update is None:
            update = OxmlElement("w:updateFields"); settings.append(update)
        update.set(qn("w:val"), "true")
        self.doc.save(path)


def _sampling_rows(results):
    labels = {"baseline_current": "基线", "temperature_0.0": "温度0", "temperature_0.5": "温度0.5",
              "top_p_0.7": "Top p 0.7", "top_k_20": "Top k 20", "naive_old_prompt": "旧规则",
              "current_prompt_reused": "新规则 复用基线"}
    rows = []
    for result in results:
        conditions = dict(result.get("config", {}).get("conditions", []))
        if not conditions:
            continue
        if "current_prompt_reused" in result.get("condition_metrics", {}):
            conditions["current_prompt_reused"] = next(iter(conditions.values()))
        for name, config in conditions.items():
            metrics = result.get("condition_metrics", {}).get(name, result["metrics"])
            label = result.get("delivery_label", "实验") + " " + labels.get(name, name)
            rows.append([label, metrics.get("case_count", "未记录"), config.get("temperature", "未记录"),
                         config.get("top_p", "未记录"), config.get("top_k_sampling", "未记录"),
                         f"{metrics['latency_mean_ms']:.1f}" if isinstance(metrics.get("latency_mean_ms"), (int, float)) else "未测量"])
    return rows


def _status_intro(writer, inputs):
    writer.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    if inputs.results:
        writer.text(f"本报告收录 {len(inputs.formal_results)} 份真实测量记录和 {sum(is_demo(r) for r in inputs.results)} 份演示模式记录。"
                    "每份记录的题数与用途单独标注；小样本功能验收不代表全量课程成绩。")
        writer.text(inputs.config.get("evaluation_note", "检索客观指标可以自动复现。答案正确性、完整性、引用质量与流畅性仍待人工评分。"))
    else:
        writer.text("系统交付源码、运行说明和实验工具。当前尚未提供可引用的真实实验结果，评测指标与优化结论保留为待完成项。")


def _experiment_identity(writer, inputs):
    rows = []
    for result in inputs.results:
        name = result.get("delivery_label", result.get("experiment_id", "实验"))
        try:
            source = str(Path(result["source_path"]).relative_to(inputs.root))
        except ValueError:
            source = result["source_path"]
        rows.append([name, result.get("delivery_scope", "实际测量"), source,
                     str(result.get("dataset_hash", "未记录"))[:12]])
    writer.table(["实验", "用途", "原始结果路径", "题集哈希前12位"], rows, [3, 2.3, 7.5, 3.2])
    writer.text("完整SHA256、配置、模型版本、索引版本和逐题输出保存在原始JSON与交付状态清单。哈希前缀仅用于表内核对。")
    if any(r.get("legacy_persisted_chunk_snapshot") for r in inputs.results):
        writer.text("M3组复用旧持久索引的359个分块，不与修订后递归分块的新变体混为同一索引。其查询向量一次真实编码后复用；表内库检索、BM25与RRF时间不含共享编码，不代表在线端到端响应时间。")
    if any(r.get("human_review_pending") for r in inputs.results):
        writer.text("原文锚点由程序核对和映射，属于来源检查草案标注；人工题集共识与答案质量评分仍待完成。")


def _validation(writer, inputs):
    rows = inputs.config.get("validation_summary", [])
    if rows:
        writer.heading("运行与自动验证证据")
        writer.table(["验证", "实际结果", "证据文件"],
                     [[row.get("name", "验证"), row.get("result", "未记录"), row.get("source", "未记录")] for row in rows], [4, 4, 8])
    writer.text("自动测试验证代码行为与故障恢复；人工论文答案评分与完整部署演示按各自实际记录认定。")


def _experiment_metrics(writer, inputs):
    retrieval = [r for r in inputs.results if r.get("kind") == "retrieval"]
    generation = [r for r in inputs.results if r.get("kind") == "generation"]
    parsers = [r for r in inputs.results if r.get("kind") == "pdf_parser_comparison"]
    databases = [r for r in inputs.results if r.get("kind") == "vector_database_comparison"]
    parameters = [r for r in inputs.results if r.get("kind") == "dev_parameter_ablation"]
    online = [r for r in inputs.results if r.get("kind") in {"online_functionality", "online_functional"}]
    encodings = [r for r in inputs.results if r.get("kind") == "embedding_encoding"]
    if parsers:
        writer.heading("PDF解析测量", 2)
        writer.table(["解析器", "论文数", "解析用时ms", "锚点保留率", "匹配锚点"],
                     [[row["parser"], row["paper_count"], f"{row['parse_ms']:.1f}", f"{row['anchor_retention']:.4f}", f"{row['matched_anchors']}/{row['total_anchors']}"]
                      for result in parsers for row in result.get("summary", [])], [3, 2, 4, 3.5, 3.5])
        writer.text("锚点保留率仅检查标注原文摘录是否仍可定位，不能证明双栏顺序、表格结构或公式语义全部正确。解析实验没有调用语言模型。")
    if databases:
        writer.heading("向量数据库实测", 2)
        writer.table(["后端", "题数", "块数", "Hit@5", "平均查询ms", "重启与增删"],
            [[row["backend"], row["query_count"], row["chunk_count"], f"{row['hit_at_5']:.4f}", f"{row['latency_mean_ms']:.3f}", "已实测" if row.get("delete_probe_passed") else "未通过"]
             for result in databases for row in result.get("records", [])], [2.5, 1.5, 1.5, 2.5, 4, 4])
        writer.text("两后端使用相同359块和同一组真实M3向量，仅查询及持久化操作参与计时，不重复模型编码。此规模下的时间不推断大型数据库吞吐。")
    if encodings:
        writer.heading("真实Embedding编码测量", 2)
        writer.table(["模型", "文档文本数", "文档编码秒", "查询条数", "查询编码秒"],
            [[r.get("embedding_model", "未记录"), r["metrics"].get("unique_document_text_count", "未记录"),
              f"{r['metrics']['document_encoding_total_ms']/1000:.2f}", r["metrics"].get("query_count", "未记录"),
              f"{r['metrics']['query_encoding_total_ms']/1000:.3f}"] for r in encodings], [5.5, 2.5, 3, 2, 3])
        writer.text("编码确实由模型执行，分块变体复用唯一文本的向量。不同模型若文档文本数或窗口总量不同，不公平比较文档批量TPS；只有相同20条查询编码可直接比较该组速度。原生token-ID完整分窗保留尾部，短文本数值一致性检查不能代替语义质量评测。")
    if parameters:
        writer.heading("开发集融合与截断消融", 2)
        writer.table(["RRF k", "指标Top K", "题数", "Hit", "Recall", "MRR@20"],
            [[row["rrf_k"], row["metric_k"], row["query_count"], f"{row['hit']:.4f}", f"{row['recall']:.4f}", f"{row['mrr_at_20']:.4f}"]
             for result in parameters for row in result.get("summary", [])], [2.2, 2.7, 1.7, 3, 3, 3.4])
        writer.text("该组复用开发集真实向量与BM25排名，仅改变确定性融合及指标截断，没有新增模型调用。测试集继续使用候选20、RRF60的固定配置。")
    if retrieval:
        writer.heading("检索与分块测量", 2)
        writer.table(["实验", "题数", "Hit@5", "MRR", "Recall@5", "检索阶段ms"],
            [[r.get("delivery_label", "实验"), r["metrics"].get("case_count", "未记录"),
              *[f"{r['metrics'][k]:.4f}" if isinstance(r["metrics"].get(k), (int, float)) else "未测量" for k in ("hit_at_5", "mrr", "recall_at_5", "latency_mean_ms")]] for r in retrieval], [4.7, 1.5, 2.1, 2, 2.5, 3.2])
    if generation:
        writer.heading("回答与Agent测量", 2)
        writer.table(["实验", "题数", "协议完成率", "工具成功率", "用时ms"],
            [[r.get("delivery_label", "实验"), r["metrics"].get("case_count", "未记录"),
              *[f"{r['metrics'][k]:.4f}" if isinstance(r["metrics"].get(k), (int, float)) else "未测量" for k in ("generation_protocol_success_rate", "tool_success_rate", "latency_mean_ms")]] for r in generation], [4.5, 1.5, 3, 3, 4])
    if online:
        writer.heading("真实线上功能样本", 2)
        writer.table(["样本数", "平均端到端秒", "输入Token总计", "输出Token总计", "平均步骤"],
            [[r["metrics"].get("sample_count", "未记录"), f"{r['metrics']['latency_mean_ms']/1000:.2f}", r["metrics"].get("input_tokens_total", "未记录"), r["metrics"].get("output_tokens_total", "未记录"), r["metrics"].get("average_agent_steps", "未记录")] for r in online], [2, 4, 4, 4, 2])
        writer.text("这是10次真实线上诊断样本，问题、Prompt与配置有变化，也包含降级结果。它们记录实际在线耗时和Token，不能当作统一任务的性能基准或因果优化结论。生成参数实验的gen-only时间另列，不能冒充在线端到端延迟。")
    if not retrieval and not generation and not parsers and not databases and not parameters and not online and not encodings:
        writer.text("当前没有可引用的实验结果。")
    writer.text("表中完成率反映请求执行状态，工具成功率反映调用执行状态；两者均不能代替答案语义正确率。开发集与测试集结论分别解读，全部指标见原始结果。")


def _screens(writer, inputs):
    if not inputs.screenshots:
        writer.text("真实运行截图尚未提供。交付前需补齐论文库、科研问答与证据定位、论文对比、实验与轨迹页面中的 3 至 5 张截图。")
    for i, shot in enumerate(inputs.screenshots[:5], 1):
        writer.image(Path(shot["path"]), f"图 {i} {shot.get('title', '系统运行界面')}")
        if shot.get("description"):
            writer.text(shot["description"])


def _bad_case_rows(inputs):
    if inputs.bad_cases:
        return [[record.get(key, "待填写") for key in
                 ("case_id", "type", "observed_error", "root_cause", "improvement", "after_result")]
                for record in inputs.bad_cases]
    return [[name, "尚未整理", "待实际实验", "待分析", "尚未记录改进", "尚未重跑"]
            for name in ("检索", "生成", "引用", "Agent")]


def build_documents(inputs: DeliveryInputs, architecture: Path | None = None, retrieval_chart: Path | None = None) -> dict[str, str]:
    inputs.output.mkdir(parents=True, exist_ok=True)
    template = Path(inputs.config.get("course_template", inputs.root / "docs/course/项目交付模板.docx"))
    if not template.is_absolute():
        template = inputs.root / template
    if not template.is_file():
        raise ValueError("课程报告需保留项目交付模板，当前路径不存在")
    artifacts = {}
    course = Writer("南京农业大学课程考核情况表", template)
    for label, value in [("学年学期", inputs.project.get("term", "待填写")), ("学院", inputs.project.get("college", "待填写")),
                         ("课程名称", inputs.project.get("course", "待填写")), ("学分", inputs.project.get("credits", "待填写")),
                         ("考核方式", "课程报告"), ("考核内容", inputs.title), ("成绩", "待教师填写"),
                         ("教师签名", "待教师填写"), ("截止日期（用户提供）", inputs.project.get("deadline", "待填写")),
                         ("实际提交日期", inputs.project.get("submission_date", "待填写"))]:
        course.text(f"{label}：{value}")
    course.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    course.doc.add_page_break()
    course.title("课程实践报告")
    course.text(inputs.title, bold=True)
    members = inputs.project.get("members") or [{"name": "待填写", "student_id": "待填写", "class_name": "待填写"}]
    course.table(["成员姓名", "班级", "学号"], [[member.get(key, "待填写") for key in ("name", "class_name", "student_id")] for member in members], [3, 5, 6])
    course.doc.add_page_break()
    course.heading("目录")
    for section in COURSE_SECTIONS:
        course.text(section)
    course.text("附录包括分块实验 Prompt 实验及 Bad Case 记录，分别保存在实验与BadCase附件文档。")
    course.doc.add_page_break()
    course.heading(COURSE_SECTIONS[0])
    course.heading("1.1 项目背景", 2)
    course.text("科研论文分散在不同文档中，查找方法、核对实验条件和比较研究结果需要反复定位原文。"
                "本项目构建通用论文知识库，结合检索增强生成与多工具调度，提供文档导入、科研问答、"
                "来源核对、论文对比和实验记录功能。课程评测选定具体研究方向，以真实论文与来源核对题目检验系统。")
    course.heading("1.2 项目目标", 2)
    course.text("实现可运行、可追踪、可复跑的科研助理。回答引用服务器生成的证据编号，保留原件与页码或段落位置；"
                "实验记录配置、语料和模型版本，使失败案例与改进结论可以核查。")
    _status_intro(course, inputs)
    course.heading("1.3 技术栈概览", 2)
    course.table(["层次", "代码与技术"], [["展示", "Streamlit 与 HTTP SSE"], ["服务", "FastAPI 与 SQLite"],
                 ["文档", "PyMuPDF pdfplumber PyPDF2 python-docx"], ["检索", "Embedding BM25 手写RRF 重排序"],
                 ["生成与调度", "Ollama RAG 手写ReAct 多工具与记忆"]], [3, 11])
    course.heading(COURSE_SECTIONS[1])
    course.table(["姓名", "角色", "负责模块", "贡献比例"], [[m.get("name", "待填写"), m.get("role", "待填写"),
                  m.get("modules", "待填写"), f"{m['contribution_percent']}%" if m.get("contribution_percent") is not None else "待填写"] for m in members], [2.5, 3, 5.5, 3])
    course.text("成员身份与40%、30%、30%贡献比例采用团队提供的信息，总和100%。角色和负责模块属于拟定分工，不作为未提供的个人实践履历。")
    course.heading(COURSE_SECTIONS[2])
    for row in module_rows(inputs):
        course.heading(row[0], 2)
        course.table(["交付项", "当前状态", "验证依据"], [row], [3.5, 3.5, 7])
    _validation(course, inputs)
    course.heading(COURSE_SECTIONS[3])
    for title, description in [("4.1 文档处理与检索", "四类文档导入保留来源位置，三种PDF解析器、固定/递归/句段边界分块可切换。向量与BM25召回经手写RRF融合，保留各阶段分数供评测。"),
            ("4.2 RAG生成", "上下文按预算组织，服务器分配S1等引用编号。流式回答、语义缓存、无结果和模型错误降级分别记录来源与状态。引用审校检查编号与证据一致性，仍需人工判断事实正确性。"),
            ("4.3 Agent决策", "手写ReAct循环校验工具动作并处理Observation。独立工具有限并行，超时、异常、重复动作和步数限制均保留轨迹。会话窗口与摘要按工作区和会话隔离。"),
            ("4.4 系统集成与前端", "前端通过HTTP读取服务，不在页面进程加载模型。论文库、科研问答、论文对比和实验记录共用来源、状态与执行数据。"),
            ("4.5 评测", "检索统计Hit@5、MRR与证据级Recall@5；生成与Agent记录真实回答、工具和Token。质量评分与四类Bad Case保留人工复核。")]:
        course.heading(title, 2); course.text(description)
    if architecture:
        course.image(architecture, "系统架构及主要数据流")
    _screens(course, inputs)
    course.heading(COURSE_SECTIONS[4])
    course.text("统一模型包括Document、SourceBlock、Chunk、RetrievalHit、Citation、ToolResult与ChatRequest。工作区凭据隔离资料；同一服务协调索引、元数据和缓存版本。接口与部署说明见技术设计文档和用户手册。")
    course.heading(COURSE_SECTIONS[5])
    _status_intro(course, inputs)
    course.text("实现状态按本报告验证清单认定。真实优化前后数据见Bad Case报告，人工评分与未完成验收按系统评测报告所列边界处理。")
    course.heading(COURSE_SECTIONS[6])
    course.text("附录A 代码目录：src/data_loader、src/chunking、src/retrieval、src/generation、src/agent、src/api、src/services、src/frontend、src/evaluation、src/delivery。")
    course.text("附录B 运行环境：使用scripts/setup.ps1安装依赖，scripts/start.ps1启动本地服务，prepare_models.py明确准备本地模型。标准模式与功能演示模式分别运行并标记。")
    course.text("附录C 原始实验文件和哈希见系统评测报告及交付状态清单。分块、Prompt、Bad Case详见独立实验附件。")
    course_path = inputs.output / "课程报告.docx"; course.save(course_path); artifacts["course_report"] = str(course_path)

    evaluation = Writer("智能科研助理系统评测报告")
    evaluation.text("本报告说明论文知识库检索、生成和工具调度的测试范围、统计口径及实际记录，供课程评审与复现使用。")
    _status_intro(evaluation, inputs)
    evaluation.heading("评测集与范围")
    evaluation.text("课程要求一个具体研究方向的10至20篇论文、50题以上问答和四类题型。题目须reviewed=true并填写参考答案与原文来源。数量统计不能证明论文真实性或人工评分完整。")
    evaluation.heading("统计口径")
    evaluation.text("Hit@5表示前五条命中任一标注证据。MRR取完整返回排名中第一条相关结果名次倒数。Recall@5按覆盖的标注证据数计算，同一证据的多个分块不重复计数。异常题计零分并保留错误。")
    evaluation.text("生成质量由人工分别评正确性、完整性、引用质量和流畅性。来源覆盖不等于答案正确性，工具成功不等于任务正确完成。Token缺失保持未测量，估算值按估算呈现。")
    evaluation.heading("实际实验指标")
    if retrieval_chart:
        evaluation.image(retrieval_chart, "同一题集与语料的检索模式对比 实际题数见图中标注")
    _experiment_metrics(evaluation, inputs)
    _validation(evaluation, inputs)
    if inputs.config.get("findings"):
        evaluation.heading("结果解读")
        for finding in inputs.config["findings"]:
            evaluation.text(finding)
    evaluation.heading("来源与复现")
    _experiment_identity(evaluation, inputs)
    evaluation.heading("失败与改进")
    evaluation.text("四类Bad Case及实际改进记录见实验与BadCase附件。只有固定语料、题目和统计规则下的前后真实测量，才能作为改进证据。")
    evaluation_path = inputs.output / "系统评测报告.docx"; evaluation.save(evaluation_path); artifacts["evaluation_report"] = str(evaluation_path)

    technical = Writer("智能科研助理技术设计文档")
    technical.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    technical.text("本设计说明系统模块、数据契约、接口和部署流程，供维护与复现使用。实现验证状态见课程报告模块清单。")
    technical.heading("系统架构")
    if architecture:
        technical.image(architecture, "展示 服务 调度 检索与持久化")
    technical.text("Streamlit调用FastAPI，服务组合导入任务、检索、RAG、Agent与实验。SQLite保存元数据、会话、运行和版本；文件系统保存原件、索引与日志。")
    technical.heading("模块接口")
    technical.table(["对象或接口", "职责"], [["Document SourceBlock Chunk", "原件、来源位置、分块与索引版本"],
        ["RetrievalHit Citation", "检索阶段分数、证据编号和来源"], ["ToolResult Agent事件", "工具结果、错误、耗时和轨迹"],
        ["POST /documents 与 GET /jobs/{id}", "批量导入与阶段进度"], ["会话接口与SSE chat", "隔离历史、流式回答、引用和Token"],
        ["experiments与export", "评测队列、真实结果和JSON CSV XLSX导出"]], [6, 10])
    technical.heading("关键实现")
    technical.text("ReAct与RRF保留手写实现。来源在分块、检索、上下文与引用中传播；上下文截短同步限制证据。缓存绑定语料版本、模型、Prompt、检索参数及会话范围，数据变化后失效。")
    technical.text("导入任务按解析、分块、Embedding、提交阶段更新真实状态。模型与解析配置变化后重建索引。失败、超时和不完整流不会冒充完成请求。")
    technical.heading("八个工具与调度契约")
    technical.table(["工具", "职责"], [["knowledge_qa", "根据已选择论文的知识库证据生成回答"],
        ["paper_metadata", "读取元信息，缺失值保持为空"], ["compare_papers", "比较2至4篇论文，核验每列来源与数值依据"],
        ["extract_keywords", "从用户问题或论文提取中英文关键词与频次"], ["summarize_paper", "按背景、方法、实验、结果、局限生成带来源摘要"],
        ["current_time", "读取真实时间，可选IANA时区"], ["calculator", "有限算术与白名单数学函数"],
        ["paper_catalog_search", "按标题、作者、年份、DOI等查目录并消除歧义"]], [5, 11])
    technical.text("ReAct循环维护模型动作与Observation，调用前按工具JSON Schema校验参数。调度器只并行独立任务，限制工作线程、工具超时、请求预算和最大步数；重复动作触发终止。失败保留错误与恢复轨迹。")
    technical.heading("配置与部署")
    technical.text("config.yaml定义模型、解析、分块、召回、生成、缓存和Agent预算。RAG_字段环境变量可覆盖配置。默认本地离线模式；模型下载通过prepare_models.py明确执行。Docker材料位于docker目录。")
    technical.text("一键Docker入口：scripts/deploy_docker.ps1 -Profile standard，可选-Gpu。当前只有脚本语法和YAML静态验证，Docker Engine未运行，不宣称容器已实际启动。模型端口、独立Ollama实例与云API配置见docs/模型与API配置.md。")
    technical.heading("边界与验证")
    technical.text("工作区标识用于本地数据隔离，默认仅绑定localhost。扫描PDF需要先OCR。复杂双栏、跨页表格和公式应人工抽样检查。语义与事实正确性依赖真实语料、模型与人工评测。")
    technical_path = inputs.output / "技术设计文档.docx"; technical.save(technical_path); artifacts["technical_design"] = str(technical_path)

    manual = Writer("智能科研助理用户手册")
    manual.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    manual.text("本手册介绍本地启动、资料导入、证据问答与实验结果查看。标准模式需要已准备的Embedding、重排与Ollama模型。")
    for title, body in [("环境与启动", "在项目目录运行scripts/setup.ps1，随后运行scripts/start.ps1。首次按README准备模型；仅验证功能时显式选择demo模式。浏览器打开启动脚本给出的Streamlit地址。"),
                      ("知识库与会话", "侧栏保存当前工作区访问令牌，用于恢复并访问对应资料。论文库支持PDF、DOCX、TXT和Markdown，等待索引状态就绪后提问。失败任务可重试，已就绪文档可按当前配置重建索引。"),
                      ("科研问答", "选择会话与论文范围，配置纯向量、混合或重排序，按需启用Agent。回答生成后展开证据，查看原文与PDF物理页高亮。关键事实应回到原文核对。"),
                      ("论文对比", "选择两篇已就绪文档及比较维度。相同任务、数据集、划分和指标条件下才解释数值差异，不可比或原文未给出时保留说明。"),
                      ("实验与运行记录", "上传已核对来源的JSONL，并注明助理或人工核对状态。先用dev调参，再固定配置跑test。可依次提交三档检索实验，查看逐题排名、失败、版本与导出。历史运行回放有明确运行ID。"),
                      ("常见问题", "API连接失败时检查服务地址与启动进程。模型未就绪时检查健康状态和本地模型缓存。空库先上传资料，扫描PDF先OCR。没有完整结束标记时应查看失败记录并重试。")]:
        manual.heading(title); manual.text(body)
    _screens(manual, inputs)
    manual_path = inputs.output / "用户手册.docx"; manual.save(manual_path); artifacts["user_manual"] = str(manual_path)

    attachment = Writer("智能科研助理分块 Prompt 与Bad Case实验记录")
    _status_intro(attachment, inputs)
    attachment.heading("分块与解析实验")
    chunk_results = [r for r in inputs.results if r.get("study") in {"chunking", "pdf", "embedding"} and r.get("kind") == "retrieval"]
    attachment.table(["实验", "策略或解析", "长度", "Hit@5", "MRR"], [[r.get("delivery_label", "实验"),
            r.get("config", {}).get("chunk_strategy", r.get("config", {}).get("pdf_parser", "未记录")),
            r.get("config", {}).get("chunk_size", "未记录"),
            *[f"{r['metrics'][key]:.4f}" if isinstance(r["metrics"].get(key), (int, float)) else "未测量" for key in ("hit_at_5", "mrr")]] for r in chunk_results] or [["尚未执行", "待对比三策略与三解析器", "256 512 1024", "未测量", "未测量"]], [3, 4.8, 2.2, 2.7, 2.8])
    attachment.text("分块实验固定语料与标注证据，比较固定切分256、512、1024，以及递归和句段边界分块。实际执行组合以表格为准。解析效果还需检查双栏、表格、公式；字符数较多不等于阅读顺序更正确。")
    attachment.heading("Prompt与生成参数实验")
    prompt_results = [r for r in inputs.results if r.get("study") in {"prompt", "generation", "context"} or r.get("kind") == "generation"]
    attachment.table(["实验", "Temperature", "Top p", "Top k", "实际用时ms"],
            [[row[0], *row[2:]] for row in _sampling_rows(prompt_results)] or [["尚未执行", "待测量", "待测量", "待测量", "未测量"]], [4, 3, 2.5, 2.5, 3.5])
    attachment.text("比较上下文预算、引用约束和生成参数时，应保存Prompt版本、实际答案、模型Token与人工质量评分。只有耗时数据不能证明Prompt质量更好。")
    attachment.heading("Bad Case与优化记录")
    for row in _bad_case_rows(inputs):
        attachment.heading(str(row[0]), 2)
        for label, value in zip(("类型", "观察到的问题", "根因", "实施改进", "改进后结果"), row[1:]):
            attachment.text(f"{label}：{value}")
    attachment.text("完整对比需记录同一题目的改进前结果、具体代码变更和固定test集上的重跑文件。尚未执行的改进不得标记已完成。")
    attachment.heading("原始实验清单")
    for source in inputs.source_files:
        try:
            source_path = str(Path(source["path"]).relative_to(inputs.root))
        except ValueError:
            source_path = source["path"]
        attachment.text(source_path + "\nSHA256 " + source["sha256"])
    attachment_path = inputs.output / "实验与BadCase附件.docx"; attachment.save(attachment_path); artifacts["experiment_attachment"] = str(attachment_path)
    artifacts.update(_independent_experiment_reports(inputs))
    artifacts["template_sha256"] = digest(template)
    return artifacts


def _independent_experiment_reports(inputs):
    """Produce the three explicitly named course outputs from measured sources."""
    artifacts = {}
    chunk = Writer("分块策略对比实验报告")
    chunk.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    chunk.text("本报告比较论文语料上固定切分256、512、1024，以及递归和句段边界分块。实际执行的策略与长度组合逐行列出；开发集用于比较配置，最终固定配置测试另行记录。")
    chunk.text("分块长度采用项目统一的轻量token计量，中文字符、英文词和标点均计入预算；它与生成模型原生tokenizer的计数不同。semantic表示句段边界分块。")
    selected = [r for r in inputs.results if r.get("study") == "chunking"]
    chunk.heading("实验配置与结果")
    chunk.table(["实验", "策略", "长度", "题数", "Hit@5", "MRR"],
        [[r.get("delivery_label", "实验"), r.get("config", {}).get("chunk_strategy", "未记录"),
          r.get("config", {}).get("chunk_size", "未记录"), r["metrics"].get("case_count", "未记录"),
          *[f"{r['metrics'][k]:.4f}" if isinstance(r["metrics"].get(k), (int, float)) else "未测量" for k in ("hit_at_5", "mrr")]] for r in selected]
        or [["尚未执行", "三策略", "256 512 1024", "无运行记录", "未测量", "未测量"]], [3.3, 3.4, 1.7, 1.5, 3, 3])
    chunk.heading("归因与选择")
    chunk.text("比较时须保持论文、问题、Embedding、召回和统计口径一致；改变分块后重新映射原文证据。较长块可能覆盖更多上下文，也可能稀释检索信号，不能仅凭块数量判定效果。")
    for finding in inputs.config.get("chunk_findings", []):
        chunk.text(finding)
    chunk.heading("解析与复杂文档边界")
    chunk.text("三解析器测量、双栏顺序、跨页表格和公式处理分别核查。文字字符数或解析耗时不能代替阅读顺序与公式语义正确性。")
    parser = [r for r in inputs.results if r.get("study") == "pdf"]
    if parser:
        parse_summary = [row for r in parser if r.get("kind") == "pdf_parser_comparison" for row in r.get("summary", [])]
        if parse_summary:
            chunk.table(["解析器", "论文数", "解析用时ms", "锚点保留率"], [[row["parser"], row["paper_count"], f"{row['parse_ms']:.1f}", f"{row['anchor_retention']:.4f}"] for row in parse_summary], [4, 3, 5, 4])
            chunk.text("10篇论文分别使用3个解析器，得到30条解析记录。119个原文锚点的自动保留检查不等同于复杂排版的人工评分。")
        retrieval_parsers = [r for r in parser if r.get("kind") == "retrieval"]
        if retrieval_parsers:
            chunk.table(["实验", "解析器", "题数", "Hit@5", "MRR"], [[r.get("delivery_label", "实验"), r.get("config", {}).get("pdf_parser", "未记录"), r["metrics"].get("case_count", "未记录"), r["metrics"].get("hit_at_5", "未测量"), r["metrics"].get("mrr", "未测量")] for r in retrieval_parsers], [3.5, 4, 2, 3.2, 3.3])
    chunk.heading("复现证据")
    for r in selected + parser:
        chunk.text(r.get("delivery_label", "实验") + "：" + str(Path(r["source_path"]).relative_to(inputs.root)))
    path = inputs.output / "分块策略对比实验报告.docx"; chunk.save(path); artifacts["chunking_report"] = str(path)

    prompt = Writer("Prompt模板设计与实验记录")
    prompt.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    prompt.text("本报告说明科研问答的证据约束、上下文组织和生成参数实验。实际模型返回及客观执行指标可以复核，答案质量仍需人工评分。")
    prompt.heading("模板设计")
    prompt.text("System Prompt定义科研助理角色、证据优先规则和引用格式。用户问题与文献内容分开组织，文献中的指令作为文献内容处理。服务器在上下文构建阶段分配S1等证据编号，生成后校验编号与来源。")
    prompt.text("无检索结果或相关性不足时明确说明证据限制；多论文比较要求每篇论文均有证据。上下文先按相关性排序，再去重并按预算截断，截断后同步限制可引用证据。")
    source = inputs.root / "src/generation/context.py"
    if source.is_file():
        prompt.text("当前Prompt实现：src/generation/context.py，SHA256 " + digest(source))
        literal = ""
        for node in ast.parse(source.read_text(encoding="utf-8")).body:
            if isinstance(node, ast.Assign) and any(isinstance(target, ast.Name) and target.id == "SYSTEM_PROMPT" for target in node.targets) and isinstance(node.value, ast.Constant) and isinstance(node.value.value, str):
                literal = node.value.value
            if isinstance(node, ast.AugAssign) and isinstance(node.target, ast.Name) and node.target.id == "SYSTEM_PROMPT" and isinstance(node.value, ast.Constant) and isinstance(node.value.value, str):
                literal += node.value.value
        if literal:
            prompt.heading("当前System Prompt", 2)
            for line in literal.strip().splitlines():
                prompt.text(line)
    prompt.heading("参数实验记录")
    if (inputs.root / "reports/experiments/generation/prepared_inputs.json").is_file():
        prompt.text("生成研究固定真实论文检索的证据输入，并保存每题上下文哈希。采样参数、Prompt与上下文预算研究复用这些输入，避免重复编码和重排；其gen-only时间不包含实时检索。准备输入本身没有调用生成模型，也不算生成成绩。设计与复跑方法见docs/生成与Agent实验记录.md和scripts/run_course_generation.py。")
    chosen = [r for r in inputs.results if r.get("study") in {"prompt", "generation", "context"} or r.get("kind") == "generation"]
    prompt.table(["实验", "题数", "T", "Top p", "Top k", "用时ms"],
        _sampling_rows(chosen)
        or [["尚未执行", "无运行记录", "未测量", "未测量", "未测量", "未测量"]], [4.5, 1.5, 1.7, 2.2, 2.2, 3.5])
    prompt.heading("实验解释与人工评分")
    prompt.text("Temperature改变采样随机性，Top p与Top k分别限制概率质量和候选词数。表格逐条件列出真实配置与生成阶段用时；新Prompt复用同4题基线答案，不重复计算样本。实验应逐次改变一个因素，不用单次耗时推断答案质量改善。人工按1至5分评正确性、完整性、引用质量和流畅性。")
    for finding in inputs.config.get("prompt_findings", []):
        prompt.text(finding)
    prompt.heading("原始记录")
    for r in chosen:
        prompt.text(r.get("delivery_label", "实验") + "：" + str(Path(r["source_path"]).relative_to(inputs.root)))
    path = inputs.output / "Prompt模板设计与实验记录.docx"; prompt.save(path); artifacts["prompt_report"] = str(path)

    bad = Writer("Bad Case分析报告")
    bad.text("报告生成日期：" + inputs.project.get("report_date", "未记录"))
    bad.text("本报告按检索、生成、引用和Agent四类整理实际观察到的失败与修复。每条记录说明来源、根因、改动和复测状态。历史功能回归与固定题集的模型效果对比单独解读。")
    bad.heading("案例与修复记录")
    for record in inputs.bad_cases:
        bad.heading(str(record.get("case_id", "案例")), 2)
        for key, label in [("type", "类别"), ("observed_error", "观察结果"), ("root_cause", "根因"), ("improvement", "实施改动"), ("before_result", "改进前记录"), ("after_result", "改进后记录"), ("improvement_status", "复测状态"), ("limitation", "比较边界"), ("evidence", "验证来源")]:
            value = record.get(key)
            if value:
                bad.text(label + "：" + ("；".join(map(str, value)) if isinstance(value, list) else str(value)))
    if not inputs.bad_cases:
        bad.text("当前尚未提供经过分析的真实案例，不能宣称已完成一轮优化。")
    bad.heading("固定数据集前后对比")
    matched = [r for r in inputs.results if r.get("study") in {"before", "after", "improvement"}]
    if matched:
        bad.table(["实验", "题数", "Hit@5", "MRR", "Recall@5"], [[r.get("delivery_label", "实验"), r["metrics"].get("case_count", "未记录"), r["metrics"].get("hit_at_5", "未测量"), r["metrics"].get("mrr", "未测量"), r["metrics"].get("recall_at_5", "未测量")] for r in matched], [5, 2, 3, 3, 3])
    for finding in inputs.config.get("bad_case_findings", []):
        bad.text(finding)
    bad.text("测试集上的后续修复意味着题集已暴露；若据此继续调参，不能把该题集继续称为未见数据。功能测试通过也不自动证明所有论文问答都正确。")
    path = inputs.output / "Bad Case分析报告.docx"; bad.save(path); artifacts["bad_case_report"] = str(path)
    return artifacts

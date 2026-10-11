from __future__ import annotations

import pytest

from src.data_loader.loaders import extract_metadata, load_blocks


def test_markdown_sections_table_formula_and_provenance(tmp_path):
    file = tmp_path / "paper.md"
    file.write_text("# A Vision Transformer Study\n\nAuthors: Alice; Bob\n\nYear: 2024\n\n"
                    "## Abstract\n\nWe test ImageNet accuracy.\n\n## Keywords\n\nattention\n\n"
                    "## Results\n\n| Model | Accuracy |\n| --- | --- |\n| ViT | 88.2 |\n\n"
                    "$$ x² + y² = z² $$\n\nDOI: 10.1234/vision.2024\n", encoding="utf-8")
    blocks = load_blocks(file)
    assert all(b.page is None for b in blocks)
    assert [b.paragraph for b in blocks] == list(range(1, len(blocks) + 1))
    table = next(b for b in blocks if b.kind == "table")
    assert "88.2" in table.text and table.section == "Results"
    assert any("x²" in b.text and b.kind == "formula" for b in blocks)
    metadata = extract_metadata(blocks, file.name)
    assert metadata["title"] == "A Vision Transformer Study"
    assert metadata["authors"] == ["Alice", "Bob"]
    assert metadata["year"] == 2024
    assert metadata["doi"] == "10.1234/vision.2024"
    assert "ImageNet" in metadata["abstract"] and "88.2" not in metadata["abstract"]


def test_text_chinese_encoding_and_empty_file(tmp_path):
    file = tmp_path / "chinese.txt"
    file.write_bytes("中文科研论文与公式 E=mc²。".encode("gb18030"))
    assert "中文科研" in load_blocks(file)[0].text
    file.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="未提取到文本"):
        load_blocks(file)


def test_docx_interleaved_paragraphs_tables_and_no_invented_pages(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading("Transformer Study", level=1)
    doc.add_paragraph("The baseline uses attention.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Method", "Accuracy"
    table.cell(1, 0).text, table.cell(1, 1).text = "ViT", "88.2"
    doc.add_paragraph("Conclusion after the table.")
    file = tmp_path / "paper.docx"
    doc.save(file)
    blocks = load_blocks(file)
    assert [b.kind for b in blocks] == ["heading", "text", "table", "text"]
    assert all(b.page is None for b in blocks)
    assert "88.2" in blocks[2].text
    assert "after" in blocks[3].text


@pytest.mark.parametrize("parser", ["pymupdf", "pdfplumber", "pypdf2"])
def test_three_pdf_parsers_extract_real_page_numbers(tmp_path, parser):
    import pymupdf
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((60, 60), "Research Paper Title", fontsize=18)
    page.insert_text((60, 90), "ImageNet accuracy is 88.2 and E = mc2.")
    page = pdf.new_page()
    page.insert_text((60, 60), "Conclusion: attention improves the representation.")
    file = tmp_path / "paper.pdf"
    pdf.save(file)
    pdf.close()
    blocks = load_blocks(file, pdf_parser=parser)
    assert {b.page for b in blocks} == {1, 2}
    assert "88.2" in " ".join(b.text for b in blocks)
    assert "Conclusion" in " ".join(b.text for b in blocks)
    if parser != "pypdf2":
        assert all(b.bbox and len(b.bbox) == 4 for b in blocks)


def test_pymupdf_two_columns_read_left_before_right(tmp_path):
    import pymupdf
    pdf = pymupdf.open()
    page = pdf.new_page(width=600, height=800)
    page.insert_text((60, 60), "Full Width Title Across Both Columns", fontsize=24)
    page.insert_text((60, 110), "LEFT_FIRST alpha paragraph", fontsize=10)
    page.insert_text((330, 110), "RIGHT_FIRST gamma paragraph", fontsize=10)
    page.insert_text((60, 145), "LEFT_SECOND beta paragraph", fontsize=10)
    page.insert_text((330, 145), "RIGHT_SECOND delta paragraph", fontsize=10)
    file = tmp_path / "columns.pdf"
    pdf.save(file)
    pdf.close()
    text = "\n".join(b.text for b in load_blocks(file))
    assert text.index("LEFT_SECOND") < text.index("RIGHT_FIRST")


def test_scanned_pdf_is_explicitly_rejected(tmp_path):
    import pymupdf
    pdf = pymupdf.open()
    pdf.new_page()
    file = tmp_path / "blank.pdf"
    pdf.save(file)
    pdf.close()
    with pytest.raises(ValueError, match="OCR"):
        load_blocks(file)


def test_metadata_does_not_invent_missing_fields(tmp_path):
    file = tmp_path / "notes.txt"
    file.write_text("Attention mechanisms\n\nA study of representations.", encoding="utf-8")
    metadata = extract_metadata(load_blocks(file), file.name)
    assert metadata["authors"] == [] and metadata["year"] is None and metadata["doi"] == ""


def test_pymupdf_ruled_tables_keep_rows_and_cross_page_provenance(tmp_path):
    import pymupdf
    pdf = pymupdf.open()
    for dataset, score in [("ImageNet", "88.2"), ("CIFAR", "90.1")]:
        page = pdf.new_page(width=400, height=400)
        for y in [60, 90, 120]:
            page.draw_line((50, y), (300, y))
        for x in [50, 200, 300]:
            page.draw_line((x, 60), (x, 120))
        page.insert_text((60, 80), "Dataset")
        page.insert_text((210, 80), "Accuracy")
        page.insert_text((60, 110), dataset)
        page.insert_text((210, 110), score)
    file = tmp_path / "table.pdf"
    pdf.save(file)
    pdf.close()
    tables = [b for b in load_blocks(file) if b.kind == "table"]
    assert len(tables) == 2
    assert tables[0].page == 1 and tables[1].page == 2
    assert "ImageNet" in tables[0].text and "88.2" in tables[0].text
    assert "CIFAR" in tables[1].text and "90.1" in tables[1].text
    assert all(table.bbox == [50.0, 60.0, 300.0, 120.0] for table in tables)

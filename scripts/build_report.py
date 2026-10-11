"""Build a factual report draft from exported results; never invent conclusions.

This script creates a Markdown report and an optional editable Word draft. A
presentation outline is included for a later reviewed slide deck. Human answer
reviews and real before/after optimisation data remain explicit requirements.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description="用真实评测结果生成报告草稿、答辩提纲和Bad Case台账")
    parser.add_argument("results", nargs="+", help="各实验导出的 results.json")
    parser.add_argument("--output", default="artifacts/course-report")
    parser.add_argument("--docx", action="store_true", help="额外输出可编辑Word草稿；需人工复核版式")
    args = parser.parse_args()
    results = [json.loads(Path(path).read_text(encoding="utf-8-sig")) for path in args.results]
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    demo = any(result.get("demo") for result in results)
    lines = ["# 智能科研助理课程实践 · 评测报告草稿", "", "> 自动整理真实实验数据，结论和人工评分仍须复核。", ""]
    if demo:
        lines += ["> 当前输入含 DEMO 功能检查结果，不得当作真实模型性能写入课程结论。", ""]
    lines += ["## 实验来源", "", "| 实验 | 类型 | 模式 | 题数 | 数据集哈希 | 配置哈希 |", "|---|---|---|---:|---|---|"]
    for result in results:
        lines.append(f"| {result.get('experiment_id', '')} | {result.get('kind')} | {result.get('mode')} | {result.get('metrics', {}).get('case_count', '')} | {result.get('dataset_hash', '')[:12]} | {result.get('config_hash', '')[:12]} |")
    lines += ["", "## 实际测量结果", ""]
    for result in results:
        lines += [f"### {result.get('kind')} / {result.get('mode')} / {result.get('experiment_id', '')[:8]}", "",
                  "| 指标 | 真实测量值 |", "|---|---:|"]
        for key, value in result.get("metrics", {}).items():
            if isinstance(value, (float, int, str)) or value is None:
                lines.append(f"| {key} | {value if value is not None else '未测量'} |")
        lines += ["", f"索引版本：`{result.get('index_version', '')}`。", ""]
        if result.get("matching_rule"):
            lines += ["证据匹配规则：" + result["matching_rule"], ""]
        for note in result.get("notes", []):
            lines.append("- " + note)
        lines.append("")
    identities = {(r.get("dataset_hash"), r.get("corpus_hash"), str(r.get("split"))) for r in results}
    lines += ["## 对比解释与下一步", ""]
    if len(identities) > 1:
        lines += ["输入实验的语料、题目或划分不同，不能直接将分数差异归因于配置。请使用相同语料和评测集重跑。", ""]
    lines += ["待人工完成：", "", "1. 核对语料来自同一具体研究方向，包含 10–20 篇真实论文。",
              "2. 完成四类题型的人工答案质量评分：正确性、完整性、引用质量、流畅性。",
              "3. 分析检索、生成、引用、Agent 四类 Bad Case，并在台账记录根因。",
              "4. 至少实施一次改进，在固定测试集上记录前后真实数据。",
              "5. 补充运行环境、模型版本、PDF复杂版式抽样检查与复现说明。", ""]
    (output / "评测报告草稿.md").write_text("\n".join(lines), encoding="utf-8")
    with (output / "BadCase台账.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        fields = ["experiment_id", "case_id", "question", "bad_case_type", "observed_error", "root_cause",
                  "improvement", "before_result", "after_result", "reviewer", "status"]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for result in results:
            for row in result.get("records", []):
                if row.get("error") or row.get("hit_at_5", 1) < 1:
                    writer.writerow({"experiment_id": result.get("experiment_id"), "case_id": row["case_id"],
                                     "question": row["question"], "bad_case_type": "检索" if result.get("kind") == "retrieval" else "待人工分类",
                                     "observed_error": row.get("error") or "Top-5未命中标注证据", "status": "待分析"})
    outline = """# 答辩演示提纲

1. 问题与用户场景：科研论文阅读、来源核对和多论文比较。
2. 课程要求与实现对应：文档、检索、RAG、手写Agent、会话、可观测性。
3. 架构和数据流：从原件到分块、索引、工具和引用。
4. 现场演示：导入 → 带引用问答 → 原页高亮 → 论文对比 → 运行回放。
5. 三档检索对比：只使用固定语料与标注题目的真实图表。
6. 参数与消融：PDF、分块、Embedding、生成参数、路由、并行、缓存。
7. Bad Case与实际改进：展示原错误、根因、改动和前后数据。
8. 自选扩展：证据审校与可复现实验工作台。
9. 部署、复现与限制：离线模型准备、启动、数据范围和资源实测。

每页应填写实际数据与真实界面截图；不将DEMO结果视为学术性能结论。
"""
    (output / "答辩提纲.md").write_text(outline, encoding="utf-8")
    if args.docx:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        document = Document()
        document.styles["Normal"].font.name = "Microsoft YaHei"
        document.styles["Normal"].font.size = Pt(10.5)
        document.styles["Normal"]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        for line in lines:
            if line.startswith("### "):
                document.add_heading(line[4:], level=2)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=1)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=0)
            elif line.startswith("|---"):
                continue
            elif line.strip():
                document.add_paragraph(line)
        section = document.sections[0]
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        section.footer.paragraphs[0]._p.append(field)
        document.save(output / "评测报告草稿.docx")
        print("Word仅为可编辑草稿，交付前请渲染检查版式并完成人工核对。")
    print(f"输出目录：{output.resolve()}")


if __name__ == "__main__":
    main()
